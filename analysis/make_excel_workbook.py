#!/usr/bin/env python3
"""Builds a ready-to-open Excel file with all your shots overlaid (like shot_snippets.png).

Give it one or more sessions (or a folder: every session in it is used). Shots from
all of them are combined, so you can overlay hundreds of shots:

    python make_excel_workbook.py ../recordings/session_20261030_163316_raw.zip
    python make_excel_workbook.py ../recordings --out ../all_my_shots.xlsx

How the chart works: Excel allows at most 255 lines (series) per chart, so instead of one
series per shot, every shot goes into ONE long column with a blank row between shots.
Excel breaks the line at each blank row, so each shot is still drawn as its own line.
Makes and misses are in separate columns, so they get separate colors.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import detector_proto as dp
import wristshot_io as wio

WINDOW_S = 1.5   # seconds before and after each release
MEASURES = [("angle_deg", "Forearm angle (°), negative = arm up"),
            ("gyro", "Wrist rotation speed (rad/s)"),
            ("accel", "Acceleration (m/s²)")]
GROUPS = [("make", "2A78D6"), ("miss", "EB6834"), ("other", "8A8985")]  # line colors (hex)
LINE_WIDTH_EMU = 6350  # 0.5 pt: thin lines so hundreds of shots stay readable
CHART_WIDTH_CM, CHART_HEIGHT_CM = 26, 9
FEATURE_COLUMNS = ["t_release", "peak_gyro", "release_accel", "hold_s", "min_pitch_deg"]


def session_paths(args: list[str]) -> list[wio.SessionPaths]:
    """Every session named on the command line; a folder means all sessions in it."""
    out = []
    for a in args:
        p = Path(a)
        if p.is_dir():
            raws = sorted({r.name.rsplit("_raw", 1)[0] for r in [*p.glob("session_*_raw.csv"), *p.glob("session_*_raw.zip")]})
            out += [wio.find_session(str(p / r)) for r in raws]
        else:
            out.append(wio.find_session(a))
    return out


def shots_of(paths: wio.SessionPaths, p: dp.DetectorParams) -> tuple[pd.DataFrame, pd.DataFrame]:
    """(raw data, one row per shot). Uses your labels if the session has them, else detected shots."""
    raw = wio.load_raw(paths.raw)
    events = wio.load_events(paths.events, raw.attrs["start_ns"])
    sig = dp.Signals.from_raw(raw, p)
    if len(events):
        shots, _ = dp.split_by_labels(sig, events, p, dp.MatchWindow())
    else:
        det = dp.detect(sig, p)
        shots = det[det["accepted"]].assign(label="other")
    shots = shots.assign(session=paths.prefix)
    shots["group"] = shots["label"].where(shots["label"].isin(["make", "miss"]), "other")
    return raw, shots


def shot_lines(raw: pd.DataFrame, shots: pd.DataFrame, first_no: int) -> list[list]:
    """Rows for the 'Shot lines' sheet: the window around each release, then a blank row."""
    t = raw["t"].to_numpy()
    values = {"angle_deg": raw["pitch_deg"].to_numpy(), "gyro": raw["gyro_mag"].to_numpy(),
              "accel": raw["acc_mag"].to_numpy()}
    blank = [None] * (4 + len(MEASURES) * len(GROUPS))
    rows = []
    for no, s in enumerate(shots.itertuples(index=False), start=first_no):
        lo, hi = t.searchsorted(s.t_release - WINDOW_S), t.searchsorted(s.t_release + WINDOW_S, "right")
        for i in range(lo, hi):
            row = [no, s.session, s.label, round(float(t[i] - s.t_release), 3)]
            for col, _ in MEASURES:
                v = round(float(values[col][i]), 3)
                row += [v if g == s.group else None for g, _ in GROUPS]
            rows.append(row)
        rows.append(blank)  # blank row = break in the line
    return rows


def add_chart(wb: Workbook, lines_ws, n_rows: int, measure_index: int, title: str, anchor: str,
              groups_present: set[str]) -> None:
    chart = ScatterChart()
    chart.title = title
    chart.style = 13
    chart.width, chart.height = CHART_WIDTH_CM, CHART_HEIGHT_CM
    chart.x_axis.title = "seconds relative to the release (0 = release)"
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = -WINDOW_S, WINDOW_S
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    # Put the axes at the left and bottom edges instead of crossing at 0.
    chart.x_axis.crosses = "min"
    chart.y_axis.crosses = "min"
    chart.legend.position = "t"
    chart.display_blanks = "gap"
    x = Reference(lines_ws, min_col=4, min_row=2, max_row=n_rows + 1)
    for gi, (group, color) in enumerate(GROUPS):
        if group not in groups_present:
            continue
        col = 5 + measure_index * len(GROUPS) + gi
        y = Reference(lines_ws, min_col=col, min_row=2, max_row=n_rows + 1)
        series = Series(y, x, title=group)
        series.marker.symbol = "none"
        series.smooth = False
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = LINE_WIDTH_EMU
        chart.series.append(series)
    wb["Charts"].add_chart(chart, anchor)


def build(sessions: list[wio.SessionPaths], out: Path) -> None:
    p = dp.DetectorParams()
    all_rows, all_shots = [], []
    next_no = 1
    for paths in sessions:
        raw, shots = shots_of(paths, p)
        all_rows += shot_lines(raw, shots, first_no=next_no)
        next_no += len(shots)
        all_shots.append(shots)
        print(f"{paths.prefix}: {len(shots)} shots")
    shots = pd.concat(all_shots, ignore_index=True)
    counts = shots["group"].value_counts().to_dict()

    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    charts_ws = wb.create_sheet("Charts")
    lines_ws = wb.create_sheet("Shot lines")
    shots_ws = wb.create_sheet("Shots")

    header = ["shot_no", "session", "label", "offset_s"]
    for col, _ in MEASURES:
        header += [f"{col}_{g}" for g, _ in GROUPS]
    lines_ws.append(header)
    for r in all_rows:
        lines_ws.append(r)

    table = shots[["session", "label"] + FEATURE_COLUMNS].copy()
    table.insert(0, "shot_no", range(1, len(table) + 1))
    shots_ws.append(list(table.columns))
    for r in table.itertuples(index=False):
        shots_ws.append([round(v, 3) if isinstance(v, float) else v for v in r])

    present = {g for g, n in counts.items() if n}
    for i, (_, title) in enumerate(MEASURES):
        add_chart(wb, lines_ws, len(all_rows), i, title, f"A{3 + i * 19}", present)
    charts_ws["A1"] = (f"{len(shots)} shots overlaid ({counts.get('make', 0)} makes = blue, "
                       f"{counts.get('miss', 0)} misses = orange"
                       + (f", {counts['other']} other = gray" if counts.get("other") else "") + ")")
    charts_ws["A1"].font = Font(bold=True, size=14)

    guide = [
        ("What this file is", True),
        (f"{len(shots)} shots from {len(sessions)} session(s): {', '.join(s.prefix for s in sessions)}.", False),
        ("Each line on the charts is one shot: 1.5 s before to 1.5 s after its release,", False),
        ("shifted so every release is at 0 on the x axis. Blue = make, orange = miss.", False),
        ("", False),
        ("The sheets", True),
        ("Charts: the three charts (arm angle, wrist rotation, acceleration).", False),
        ("Shot lines: the numbers behind the charts. One row per sample (100 per second).", False),
        ("   offset_s = seconds from that shot's release. A blank row separates one shot from the next;", False),
        ("   that blank row is what makes Excel start a new line for each shot.", False),
        ("   angle_deg_make holds the value only if that shot was a make, angle_deg_miss only if a miss.", False),
        ("   That split is how makes and misses get different colors.", False),
        ("Shots: one row per shot: when it happened (t_release, seconds into the session),", False),
        ("   flick (peak_gyro, rad/s), spike (release_accel, m/s²), follow-through (hold_s, s),", False),
        ("   and how high the arm went (min_pitch_deg, °).", False),
        ("", False),
        ("How to read it", True),
        ("Where the lines bunch tightly, your motion was the same on every shot.", False),
        ("Where they spread out, that part of your shot varies. See docs/reading_the_plots.md.", False),
    ]
    for i, (text, bold) in enumerate(guide, start=1):
        start.cell(row=i, column=1, value=text).font = Font(bold=bold, size=12 if bold else 11)
    start.column_dimensions["A"].width = 100
    for ws in (lines_ws, shots_ws):
        ws.freeze_panes = "A2"
        for c in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(c)].width = 14

    wb.active = wb.index(charts_ws)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Saved {out}: {len(shots)} shots, {len(all_rows)} rows")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("sessions", nargs="+", help="session files/prefixes, or folders (all sessions inside)")
    ap.add_argument("--out", default=str(Path(__file__).parent / "output" / "shots_overlay.xlsx"))
    args = ap.parse_args()
    build(session_paths(args.sessions), Path(args.out))


if __name__ == "__main__":
    main()
